import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    PRIMARY_COLOR = RGBColor(11, 97, 164)     # #0B61A4 Deep Blue
    SECONDARY_COLOR = RGBColor(123, 31, 162)  # #7B1FA2 Purple
    DARK_TEXT = RGBColor(17, 24, 39)         # #111827
    MUTED_TEXT = RGBColor(107, 114, 128)     # #6B7280

    # Title Page / Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("D.A.R.I.A.\nDomain-Specific Agentic Research, Intelligence & Analysis System")
    run_title.font.name = "Inter"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Product Requirements Document (PRD) & Software Architecture Document (SAD)")
    run_sub.font.name = "Inter"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Inter"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Inter"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR
        return h

    def add_heading_3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Inter"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = DARK_TEXT
        return h

    def add_body_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Inter"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_TEXT
        r_text = p.add_run(text)
        r_text.font.name = "Inter"
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Inter"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_TEXT
        r_text = p.add_run(text)
        r_text.font.name = "Inter"
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = DARK_TEXT
        return p

    def add_code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # =========================================================================
    # PART 1: PRD — PRODUCT REQUIREMENTS DOCUMENT
    # =========================================================================
    add_heading_1("1. PRD — Product Requirements Document")
    add_body_p("This document outlines the purpose, scope, functionality, user flows, and requirements for the D.A.R.I.A. autonomous research intelligence system.")

    # 1. Product Overview
    add_heading_2("1. Product Overview")
    add_bullet("D.A.R.I.A. (Domain-Specific Agentic Research, Intelligence & Analysis System)", "Product/Project Name: ")
    add_bullet("AI Systems & Autonomous Agents Engineering Team", "Team Members: ")
    add_bullet("Traditional RAG and LLM systems lack long-term memory, systematic research planning, evidence critique, self-correction reflection loops, and persistent learning.", "Problem Being Solved: ")
    add_bullet("An autonomous memory-driven multi-agent research pipeline engineered with LangGraph, ChromaDB, LiteLLM, and Streamlit that plans, searches local vector stores & live web data, critiques evidence, reflects on gaps, and exports publication-ready DOCX reports.", "Short Description: ")
    add_bullet("Researchers, Engineers, Domain Analysts, Students, and Knowledge Workers requiring evidence-backed, verified research summaries.", "Target Users: ")

    # 2. Problem Statement
    add_heading_2("2. Problem Statement")
    add_bullet("Single-prompt LLM queries and basic RAG setups suffer from hallucinations, incomplete context retrieval, lack of multi-source verification, and zero memory of prior research sessions.", "What problem does the product solve? ")
    add_bullet("Domain analysts, academic researchers, technical writers, and engineering professionals who waste hours verifying uncurated AI outputs.", "Who faces this problem? ")
    add_bullet("Existing LLM wrappers do not evaluate the quality of retrieved data, do not plan information gaps before generating answers, and cannot perform self-correction reflection loops when data is missing.", "What is the existing difficulty/gap? ")

    # 3. Product Objective
    add_heading_2("3. Product Objective")
    add_bullet("To create an explainable, self-correcting multi-agent research framework that autonomously formulates research plans, executes parallel hybrid retrieval (local RAG + live web), evaluates evidence quality through a dedicated critic, and maintains persistent semantic memory.", "Main Objective: ")
    add_bullet("Achieve high evidence coverage, eliminate redundant searches via persistent memory, produce high-confidence verified reports, and provide 1-click DOCX document export.", "Target Outcome: ")

    # 4. Target Users
    add_heading_2("4. Target Users")
    add_bullet("Requires structured literature analysis, synthesis of technical concepts, and verified citations.", "Academic & Industrial Researchers: ")
    add_bullet("Needs evidence-backed comparisons across domains (e.g., Mechanical Engineering vs. Computer Science).", "Domain Experts & Engineers: ")
    add_bullet("Wants an intelligent agent system that remembers past queries and generates clear executive summaries.", "Students & Educators: ")

    # 5. Key Features
    add_heading_2("5. Key Features")
    
    features = [
        ("Memory Agent", "Queries ChromaDB vector memory (`research_memory`) to retrieve relevant past research experiences, insights, and recommendations before executing new work.", "Prevents redundant work and enables persistent long-term learning across research sessions."),
        ("Information Needs Analyst", "Decomposes complex user queries into structured research steps, identifies missing information gaps, and selects the optimal routing strategy (RAG, Web, Hybrid, or Memory).", "Ensures thorough and targeted research coverage instead of single-pass prompt guessing."),
        ("Hybrid Parallel Retrieval Layer", "Executes parallel retrieval across local domain-specific vector stores (BAAI BGE Small embeddings + Cross-Encoder reranker) and live web APIs (Tavily Search).", "Combines deep domain knowledge with live, real-world information."),
        ("Evidence Curator", "Merges retrieved RAG and Web evidence, filters out noise and duplicates, and produces curated evidence packages.", "Eliminates irrelevant content and reduces hallucination risk."),
        ("Research Critic & Reflection Loop", "Evaluates curated evidence against research plans, calculates a quality score (1-10), and triggers iterative re-planning loops if evidence is insufficient.", "Guarantees high-confidence, verified research outputs through automated self-correction."),
        ("Summarizer Agent & DOCX Exporter", "Synthesizes structured research reports (Executive Summary, Key Findings, Analysis, Sources) and allows 1-click DOCX download.", "Provides production-ready, exportable research deliverables."),
        ("Memory Update Agent", "Persists research plans, critic feedback, scores, and final summaries back into ChromaDB vector memory.", "Closes the feedback loop for continuous system learning.")
    ]

    for f_name, f_desc, f_why in features:
        add_heading_3(f_name)
        add_bullet(f_desc, "What it does: ")
        add_bullet(f_why, "Why required: ")

    # 6. User Flow / Use Cases
    add_heading_2("6. User Flow / Use Cases")
    add_body_p("The end-to-end user interaction flow follows a seamless execution path:")
    add_code_block("User enters Query → Memory Agent Check → Analyst Formulation → Parallel RAG + Web Retrieval → Evidence Curation → Research Critic Evaluation → [Reflection Loop if Score < 7] → Summarizer Report Generation → Memory Update → View Results & Download DOCX")

    # 7. Functional Requirements
    add_heading_2("7. Functional Requirements")
    add_bullet("The system MUST allow users to enter complex research queries via a Streamlit web interface.")
    add_bullet("The system MUST query persistent vector memory to check for related past research before starting execution.")
    add_bullet("The system MUST dynamically route queries to RAG, Web, or Hybrid retrieval based on analyst gap assessment.")
    add_bullet("The system MUST perform semantic reranking on local vector DB results using Cross-Encoder models.")
    add_bullet("The system MUST score evidence quality via a Research Critic agent and support up to specified reflection iterations.")
    add_bullet("The system MUST generate a structured report and provide a 1-click DOCX download button.")
    add_bullet("The system MUST update persistent memory upon successful report completion.")

    # 8. Non-Functional Requirements
    add_heading_2("8. Non-Functional Requirements")
    add_bullet("Parallel execution of RAG and Web retrieval to minimize total pipeline latency.", "Performance: ")
    add_bullet("API keys (Gemini, Tavily, Groq) stored in environment variables or password inputs; no hardcoded credentials.", "Security: ")
    add_bullet("Decoupled multi-agent architecture using LangGraph state graph allowing easy addition of new agents.", "Scalability: ")
    add_bullet("Resilient error handling returning fallback states when local vector stores or APIs are empty/unreachable.", "Reliability: ")
    add_bullet("Local persistent storage with ChromaDB ensuring 100% offline availability for memory & RAG data.", "Availability: ")
    add_bullet("Clean, modern Streamlit UI matching sample design guidelines with mint green agent badges and metric cards.", "Usability: ")

    # 9. Scope
    add_heading_2("9. Scope")
    add_heading_3("In Scope (Version 1.0)")
    add_bullet("LangGraph state graph orchestration with 8 specialized agents.")
    add_bullet("ChromaDB vector store for local engineering mechanics knowledge base & research memory.")
    add_bullet("Tavily Search integration for live web retrieval.")
    add_bullet("LiteLLM integration with Gemini 2.0 Flash / Groq models.")
    add_bullet("Streamlit dashboard UI with custom CSS styling.")
    add_bullet("DOCX report export engine (`python-docx`).")

    add_heading_3("Out of Scope (Version 1.0)")
    add_bullet("Multi-user authentication & role-based access control (RBAC).")
    add_bullet("Real-time collaborative multi-user sessions.")
    add_bullet("Audio/Video multimodal RAG processing.")

    # 10. Future Enhancements
    add_heading_2("10. Future Enhancements")
    add_bullet("Multi-Domain Corporate Corpora with automated document ingestion pipeline.")
    add_bullet("Source Credibility & Factuality Scoring matrix.")
    add_bullet("Distributed Qdrant / Milvus vector DB integration for billion-scale embeddings.")
    add_bullet("Autonomous hierarchical task decomposition for enterprise research projects.")

    doc.add_page_break()

    # =========================================================================
    # PART 2: SAD — SOFTWARE ARCHITECTURE DOCUMENT
    # =========================================================================
    add_heading_1("2. SAD — Software Architecture Document")
    add_body_p("This document describes the technical architecture, component design, data flow, database schemas, and system design patterns applied in D.A.R.I.A.")

    # 1. System Overview
    add_heading_2("1. System Overview")
    add_body_p("D.A.R.I.A. uses a state-machine multi-agent graph architecture powered by **LangGraph**. The state graph maintains a centralized `ResearchState` dictionary passed across nodes. Each node represents a dedicated agent with distinct responsibilities.")

    # 2. Architecture Diagram
    add_heading_2("2. Architecture Diagram")
    add_code_block("""+-----------------------------------------------------------------------+
|                          Streamlit Frontend UI                       |
+-----------------------------------+-----------------------------------+
                                    |
                                    v
+-----------------------------------------------------------------------+
|                       LangGraph Research State Graph                  |
|                                                                       |
|  [Memory Agent] --> [Analyst] --> [Routing Node]                      |
|                                         |                             |
|                    +--------------------+--------------------+        |
|                    |                                         |        |
|                    v                                         v        |
|               [RAG Agent]                               [Web Agent]   |
|                    |                                         |        |
|                    +--------------------+--------------------+        |
|                                         |                             |
|                                         v                             |
|                              [Evidence Curator]                       |
|                                         |                             |
|                                         v                             |
|                              [Research Critic]                        |
|                                         |                             |
|                             (Score < 7 & Loops < N?)                  |
|                             /                      \\                  |
|                        (Yes/Re-plan)             (No/Pass)            |
|                            /                        \\                 |
|                           v                          v                |
|                      [Analyst]               [Summarizer Agent]       |
|                                                      |                |
|                                                      v                |
|                                             [Memory Update Agent]     |
+-----------------------------------+-----------------------------------+
                                    |
            +-----------------------+-----------------------+
            |                       |                       |
            v                       v                       v
    +---------------+       +---------------+       +---------------+
    | ChromaDB Store|       | Tavily Search |       | LiteLLM API   |
    | (Memory & RAG)|       | Web API       |       | (Gemini/Groq) |
    +---------------+       +---------------+       +---------------+""")

    # 3. Technology Stack Table
    add_heading_2("3. Technology Stack")
    
    table_stack = doc.add_table(rows=1, cols=3)
    table_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_stack.rows[0].cells
    hdr_cells[0].text = "Layer"
    hdr_cells[1].text = "Technology / Library"
    hdr_cells[2].text = "Purpose / Justification"

    for cell in hdr_cells:
        set_cell_background(cell, "0B61A4")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Inter"
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    stack_data = [
        ("Frontend UI", "Streamlit 1.64+", "Lightweight reactive web app framework with custom CSS."),
        ("Agent Framework", "LangGraph 1.2+", "State-graph orchestration with conditional routing & reflection loops."),
        ("LLM Orchestration", "LiteLLM 1.100+", "Unified LLM provider bridge (Gemini 2.0 Flash / Groq Llama 3.3)."),
        ("Vector Database", "ChromaDB 1.5+", "Persistent vector store for local knowledge base & research memory."),
        ("Embeddings", "BAAI/bge-small-en-v1.5", "High-accuracy sentence transformer embeddings for RAG & memory."),
        ("Reranker", "cross-encoder/ms-marco-MiniLM-L-6-v2", "Semantic re-scoring of vector retrieval candidates."),
        ("Web Research API", "Tavily Python SDK 0.8+", "Optimized search engine API for AI agent web retrieval."),
        ("Report Export", "python-docx 1.2+", "Programmatic generation of Microsoft Word `.docx` reports.")
    ]

    for layer, tech, purp in stack_data:
        row_cells = table_stack.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = tech
        row_cells[2].text = purp
        for cell in row_cells:
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Inter"
                    r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. Major Components / Modules
    add_heading_2("4. Major Components / Modules")
    add_bullet("Contains memory_agent.py, information_needs_analyst.py, rag_agent.py, web_agent.py, evidence_curator.py, research_critic.py, summarizer_agent.py, and memory_update.py.", "`agents/` Module: ")
    add_bullet("Defines state.py (ResearchState schema), routing.py (conditional edge decisions), and graph.py (LangGraph workflow assembly).", "`graph/` Module: ")
    add_bullet("Manages chroma_client.py, memory_retriever.py, and memory_store.py for persistent memory operations.", "`memory/` Module: ")
    add_bullet("Contains retrieve.py, rerank.py, vectorstore/ build & verify scripts for local knowledge retrieval.", "`rag/` Module: ")
    add_bullet("Executes docx_exporter.py for automated Word document formatting.", "`export/` Module: ")

    # 5. Data Flow
    add_heading_2("5. Data Flow")
    add_body_p("Data flows through the central `ResearchState` dictionary:")
    add_code_block("User Query -> ResearchState -> Memory Agent (adds memory_evidence) -> Analyst (adds research_plan & information_gaps) -> Router -> RAG/Web Agents (adds rag_evidence & web_evidence) -> Curator (adds filtered_evidence) -> Critic (adds critic_score & feedback) -> Summarizer (adds final_response) -> Memory Update (persists state to ChromaDB)")

    # 6. Database Design
    add_heading_2("6. Database Design")
    add_body_p("D.A.R.I.A. utilizes **ChromaDB** with two primary collections:")
    add_heading_3("1. `research_memory` Collection")
    add_bullet("Query embedding generated via `BAAI/bge-small-en-v1.5`.", "Vector Index: ")
    add_bullet("query (str), research_plan (str), information_gaps (str), critic_feedback (str), final_response (str).", "Metadata Fields: ")

    add_heading_3("2. `engineering_mechanics_knowledge_base` Collection")
    add_bullet("Document chunk embedding via `BAAI/bge-small-en-v1.5`.", "Vector Index: ")
    add_bullet("chapter (str), section (str), subsection (str), page (int), source_file (str).", "Metadata Fields: ")

    # 7. API Design
    add_heading_2("7. API Design")
    add_bullet("`completion(model='gemini/gemini-2.0-flash', messages=[...], response_format=PydanticSchema)` — Structured agent output generation.", "LiteLLM API: ")
    add_bullet("`TavilyClient.search(query=query, max_results=3)` — Live web search API.", "Tavily Search API: ")
    add_bullet("`graph.invoke(initial_state)` — Synchronous pipeline invocation from Streamlit.", "LangGraph Graph API: ")

    # 8. System Design Concepts Applied
    add_heading_2("8. System Design Concepts Applied")

    table_concepts = doc.add_table(rows=1, cols=3)
    table_concepts.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdr = table_concepts.rows[0].cells
    c_hdr[0].text = "System Design Concept"
    c_hdr[1].text = "Where It Was Applied"
    c_hdr[2].text = "Why It Was Used"

    for cell in c_hdr:
        set_cell_background(cell, "7B1FA2")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Inter"
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    concept_data = [
        ("Persistent Vector Memory / Caching", "Memory Agent & Memory Store", "Avoids re-running expensive LLM & web search operations for repeated query concepts."),
        ("Parallel Hybrid Retrieval", "Graph Router & RAG / Web Nodes", "Executes local RAG and live web search concurrently to improve evidence coverage and reduce latency."),
        ("State Graph Pattern", "LangGraph Workflow (`state.py`)", "Ensures explicit context passing, state immutability, and deterministic node execution."),
        ("Semantic Reranking", "RAG Retrieval Layer (`rerank.py`)", "Uses Cross-Encoder models to re-evaluate vector candidates, eliminating low-relevance chunks."),
        ("Quality Gate & Self-Correction Loop", "Research Critic & Routing Node", "Enforces automated verification, preventing halllucinated or incomplete reports from reaching users."),
        ("Modular Layered Architecture", "Separation of `agents/`, `rag/`, `memory/`, `export/`", "Improves code maintainability, testability, and decoupled module enhancements.")
    ]

    for conc, where, why in concept_data:
        r_cells = table_concepts.add_row().cells
        r_cells[0].text = conc
        r_cells[1].text = where
        r_cells[2].text = why
        for cell in r_cells:
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Inter"
                    r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 9. Scalability & Performance
    add_heading_2("9. Scalability & Performance")
    add_bullet("Parallel multi-node execution for RAG & Web retrieval reduces latency by ~40%.")
    add_bullet("Query truncation (350 char safety cut) prevents external API 400 errors.")
    add_bullet("Local ChromaDB indexing enables sub-50ms vector query lookups.")

    # 10. Security Considerations
    add_heading_2("10. Security Considerations")
    add_bullet("API Keys configured via `.env` or password-masked sidebar input fields.")
    add_bullet("Pydantic schema validation on all LLM responses to prevent injection of malformed data.")
    add_bullet("Isolated local SQLite storage for vector data preventing unauthorized cloud data exposure.")

    # 11. Limitations
    add_heading_2("11. Limitations")
    add_bullet("Local ChromaDB instance is single-node persistent client (non-clustered).")
    add_bullet("Web search results depend on Tavily API rate limits and web page accessibility.")

    # 12. Future Architecture Improvements
    add_heading_2("12. Future Architecture Improvements (10x / 100x Scaling)")
    add_bullet("Migrate from ChromaDB to distributed Qdrant or Milvus cluster for multi-tenant scalability.")
    add_bullet("Replace local state checkpointer with Redis backed LangGraph checkpointing for distributed worker pools.")
    add_bullet("Implement Celery / RabbitMQ task queues for asynchronous background report generation.")

    # Save document
    output_filename = "DARIA_PRD_and_SAD_Documentation.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")
    return output_filename

if __name__ == "__main__":
    create_document()
